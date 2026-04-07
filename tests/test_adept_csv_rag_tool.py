from __future__ import annotations

import importlib
import sys
from pathlib import Path
from types import SimpleNamespace
from typing import Any
from unittest.mock import Mock

import pandas as pd
import pytest
from docx import Document


def _import_real_csv_rag_tool() -> Any:
    module_name = "agentic_framework_pkg.mcp_server.tools.csv_rag_tool"
    existing = sys.modules.get(module_name)
    if isinstance(existing, Mock):
        sys.modules.pop(module_name, None)
    return importlib.import_module(module_name)


csv_rag_tool = _import_real_csv_rag_tool()


class FakeCtx:
    def __init__(self) -> None:
        self.infos: list[str] = []
        self.errors: list[str] = []
        self.warnings: list[str] = []

    async def info(self, message: str) -> None:
        self.infos.append(message)

    async def error(self, message: str) -> None:
        self.errors.append(message)

    async def warning(self, message: str) -> None:
        self.warnings.append(message)


class FakeMCP:
    def __init__(self) -> None:
        self.registered: list[str] = []

    def tool(self):
        def decorator(func: Any) -> Any:
            self.registered.append(func.__name__)
            return func

        return decorator


class FakeLLMClient:
    def __init__(self, embeddings: list[list[float]] | None = None, response_text: str = "answer") -> None:
        self.embeddings = embeddings if embeddings is not None else [[0.1, 0.2]]
        self.response_text = response_text
        self.embedding_calls: list[list[str]] = []
        self.response_calls: list[dict[str, Any]] = []

    async def acreate_embedding(self, input_texts: list[str]) -> list[list[float]]:
        self.embedding_calls.append(input_texts)
        return self.embeddings

    async def agenerate_response(self, **kwargs: Any) -> Any:
        self.response_calls.append(kwargs)
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=self.response_text))])


class RaisingLLMClient:
    async def acreate_embedding(self, input_texts: list[str]) -> list[list[float]]:
        del input_texts
        raise csv_rag_tool.LLMServiceError("embed failed")

    async def agenerate_response(self, **kwargs: Any) -> Any:
        del kwargs
        raise csv_rag_tool.LLMServiceError("answer failed")


class FakeVectorStoreManager:
    def __init__(self) -> None:
        self.add_calls: list[dict[str, Any]] = []
        self.query_result: dict[str, Any] = {"documents": [["chunk 1", "chunk 2"]]}
        self.query_calls: list[dict[str, Any]] = []

    def add_documents(self, **kwargs: Any) -> None:
        self.add_calls.append(kwargs)

    def query_collection(self, **kwargs: Any) -> dict[str, Any]:
        self.query_calls.append(kwargs)
        return self.query_result


@pytest.fixture(autouse=True)
def reset_csv_rag_globals(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(csv_rag_tool, "_llm_client_instance", None)
    monkeypatch.setattr(csv_rag_tool, "_nvidia_async_client", None)
    monkeypatch.setattr(csv_rag_tool, "_nvidia_api_key", None)
    monkeypatch.setattr(csv_rag_tool, "_nvidia_base_url", "https://integrate.api.nvidia.com/v1")
    monkeypatch.setattr(csv_rag_tool, "_nvidia_model_name", "nvidia/llama-3.1-nemotron-nano-vl-8b-v1")


@pytest.mark.unit
def test_text_readers_and_chunking(tmp_path: Path) -> None:
    csv_path = tmp_path / "sample.csv"
    csv_path.write_text("A,B\n1,2\n3,4\n", encoding="utf-8")
    assert csv_rag_tool._read_csv_to_texts(str(csv_path)) == ["A: 1, B: 2", "A: 3, B: 4"]

    xlsx_path = tmp_path / "sample.xlsx"
    pd.DataFrame([{"A": 1, "B": 2}]).to_excel(xlsx_path, index=False)
    assert any("Sheet:" in row for row in csv_rag_tool._read_excel_to_texts(str(xlsx_path)))

    txt_path = tmp_path / "sample.txt"
    txt_path.write_text("hello world\n" * 20, encoding="utf-8")
    assert csv_rag_tool._read_txt_to_texts(str(txt_path))

    tex_path = tmp_path / "sample.tex"
    tex_path.write_text("\\section{Hi}\ncontent here", encoding="utf-8")
    assert csv_rag_tool._read_latex_to_texts(str(tex_path))

    docx_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Paragraph one")
    doc.add_paragraph("Paragraph two")
    doc.save(str(docx_path))
    assert csv_rag_tool._read_docx_to_texts(str(docx_path)) == ["Paragraph one\nParagraph two"]

    assert csv_rag_tool._split_text_into_chunks("", 10, 2) == []
    assert csv_rag_tool._split_text_into_chunks("abcdefghij", 4, 1) == ["abcd", "defg", "ghij", "j"]


@pytest.mark.unit
@pytest.mark.asyncio
async def test_extract_text_from_image_requires_nvidia_client() -> None:
    with pytest.raises(csv_rag_tool.LLMServiceError, match="not initialized"):
        await csv_rag_tool._extract_text_from_image_content_with_nvidia(b"img", "image/png", "prompt")


@pytest.mark.unit
@pytest.mark.asyncio
async def test_extract_text_from_image_success_and_empty(monkeypatch: pytest.MonkeyPatch) -> None:
    class FakeCompletions:
        def __init__(self, content: str | None) -> None:
            self.content = content

        async def create(self, **kwargs: Any) -> Any:
            del kwargs
            message = SimpleNamespace(content=self.content)
            return SimpleNamespace(choices=[SimpleNamespace(message=message)])

    csv_rag_tool._nvidia_async_client = SimpleNamespace(chat=SimpleNamespace(completions=FakeCompletions("ocr text")))
    text = await csv_rag_tool._extract_text_from_image_content_with_nvidia(b"img", "image/png", "prompt")
    assert text == "ocr text"

    csv_rag_tool._nvidia_async_client = SimpleNamespace(chat=SimpleNamespace(completions=FakeCompletions(None)))
    assert await csv_rag_tool._extract_text_from_image_content_with_nvidia(b"img", "image/png", "prompt") == ""


@pytest.mark.unit
@pytest.mark.asyncio
async def test_process_uploaded_file_basic_error_paths(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    ctx = FakeCtx()
    missing = await csv_rag_tool.process_uploaded_file(str(tmp_path / "missing.csv"), "missing.csv", ctx)
    assert "error" in missing

    csv_rag_tool._llm_client_instance = FakeLLMClient()

    async def fake_session(ctx: Any, override_session_id: str | None = None) -> str:
        del ctx, override_session_id
        return "session-1"

    monkeypatch.setattr(csv_rag_tool, "ensure_session_initialized", fake_session)
    unsupported_path = tmp_path / "sample.xyz"
    unsupported_path.write_text("data", encoding="utf-8")
    result = await csv_rag_tool.process_uploaded_file(str(unsupported_path), unsupported_path.name, ctx)
    assert result["error"].startswith("Unsupported file type")


@pytest.mark.unit
@pytest.mark.asyncio
async def test_process_uploaded_file_handles_empty_text_and_read_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    ctx = FakeCtx()
    csv_rag_tool._llm_client_instance = FakeLLMClient()

    async def fake_session(ctx: Any, override_session_id: str | None = None) -> str:
        del ctx, override_session_id
        return "session-1"

    monkeypatch.setattr(csv_rag_tool, "ensure_session_initialized", fake_session)

    empty_txt = tmp_path / "empty.txt"
    empty_txt.write_text("", encoding="utf-8")
    empty_result = await csv_rag_tool.process_uploaded_file(str(empty_txt), empty_txt.name, ctx)
    assert "No text content could be extracted" in empty_result["error"]

    bad_csv = tmp_path / "bad.csv"
    bad_csv.write_text("x", encoding="utf-8")
    monkeypatch.setattr(csv_rag_tool, "_read_csv_to_texts", lambda path: (_ for _ in ()).throw(ValueError("bad read")))
    error_result = await csv_rag_tool.process_uploaded_file(str(bad_csv), bad_csv.name, ctx)
    assert error_result["error"].startswith("Failed to read or process file")


@pytest.mark.unit
@pytest.mark.asyncio
async def test_process_uploaded_file_success_updates_session_and_removes_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    ctx = FakeCtx()
    csv_rag_tool._llm_client_instance = FakeLLMClient(embeddings=[[0.1, 0.2], [0.3, 0.4]])
    fake_vsm = FakeVectorStoreManager()
    updated_contexts: list[dict[str, Any]] = []

    async def fake_session(ctx: Any, override_session_id: str | None = None) -> str:
        del ctx, override_session_id
        return "session-1"

    async def fake_get_session_context(session_id: str) -> dict[str, Any]:
        assert session_id == "session-1"
        return {"uploaded_csv_files": {}}

    async def fake_update_session_context(session_id: str, new_context: dict[str, Any]) -> None:
        updated_contexts.append({"session_id": session_id, **new_context})

    monkeypatch.setattr(csv_rag_tool, "ensure_session_initialized", fake_session)
    monkeypatch.setattr(csv_rag_tool, "get_session_context", fake_get_session_context)
    monkeypatch.setattr(csv_rag_tool, "update_session_context", fake_update_session_context)
    monkeypatch.setattr(csv_rag_tool, "VectorStoreManager", lambda: fake_vsm)
    monkeypatch.setattr(csv_rag_tool.uuid, "uuid4", lambda: "uuid-fixed")

    csv_path = tmp_path / "sample.csv"
    csv_path.write_text("A,B\n1,2\n3,4\n", encoding="utf-8")

    result = await csv_rag_tool.process_uploaded_file(str(csv_path), csv_path.name, ctx, mcp_session_id="override")

    assert result["message"] == "File processed and indexed."
    assert result["file_id"] == "uuid-fixed"
    assert fake_vsm.add_calls
    assert updated_contexts[0]["uploaded_csv_files"]["uuid-fixed"]["status"] == "indexed"
    assert not csv_path.exists()


@pytest.mark.unit
@pytest.mark.asyncio
async def test_process_uploaded_file_pdf_and_image_fallback_paths(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    ctx = FakeCtx()
    csv_rag_tool._llm_client_instance = FakeLLMClient()

    async def fake_session(ctx: Any, override_session_id: str | None = None) -> str:
        del ctx, override_session_id
        return "session-1"

    async def fake_get_session_context(session_id: str) -> dict[str, Any]:
        return {"uploaded_csv_files": {}}

    async def fake_update_session_context(session_id: str, new_context: dict[str, Any]) -> None:
        return None

    monkeypatch.setattr(csv_rag_tool, "ensure_session_initialized", fake_session)
    monkeypatch.setattr(csv_rag_tool, "get_session_context", fake_get_session_context)
    monkeypatch.setattr(csv_rag_tool, "update_session_context", fake_update_session_context)
    monkeypatch.setattr(csv_rag_tool, "VectorStoreManager", lambda: FakeVectorStoreManager())
    monkeypatch.setattr(csv_rag_tool.uuid, "uuid4", lambda: "uuid-fixed")

    pdf_path = tmp_path / "sample.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")
    csv_rag_tool._nvidia_async_client = object()
    monkeypatch.setattr(
        csv_rag_tool,
        "_get_fitz_module",
        lambda: SimpleNamespace(open=lambda path: (_ for _ in ()).throw(RuntimeError("pdf fail"))),
    )
    monkeypatch.setattr(csv_rag_tool, "_read_pdf_to_texts", lambda path: ["fallback pdf text"])
    pdf_result = await csv_rag_tool.process_uploaded_file(str(pdf_path), pdf_path.name, ctx)
    assert pdf_result["collection_name"] == "session-1_uuid-fixed"

    image_path = tmp_path / "image.png"
    image_path.write_bytes(b"pngdata")
    monkeypatch.setattr(
        csv_rag_tool,
        "_extract_text_from_image_content_with_nvidia",
        lambda image_bytes, image_mime_type, prompt_text: (_ for _ in ()).throw(
            csv_rag_tool.LLMServiceError("ocr fail")
        ),
    )
    image_result = await csv_rag_tool.process_uploaded_file(str(image_path), image_path.name, ctx)
    assert "No text content could be extracted" in image_result["error"]


@pytest.mark.unit
@pytest.mark.asyncio
async def test_query_file_content_error_paths(monkeypatch: pytest.MonkeyPatch) -> None:
    ctx = FakeCtx()
    assert "LLM client not available" in await csv_rag_tool.query_file_content("q", "file", ctx)

    csv_rag_tool._llm_client_instance = FakeLLMClient(embeddings=[])

    async def fake_session(ctx: Any, override_session_id: str | None = None) -> str:
        del ctx, override_session_id
        return "session-1"

    async def empty_context(session_id: str) -> dict[str, Any]:
        del session_id
        return {"uploaded_csv_files": {}}

    async def context_without_collection(session_id: str) -> dict[str, Any]:
        del session_id
        return {"uploaded_csv_files": {"f1": {"original_filename": "doc.txt"}}}

    async def context_with_collection(session_id: str) -> dict[str, Any]:
        del session_id
        return {"uploaded_csv_files": {"f1": {"original_filename": "doc.txt", "collection_name": "col"}}}

    monkeypatch.setattr(csv_rag_tool, "ensure_session_initialized", fake_session)
    monkeypatch.setattr(csv_rag_tool, "get_session_context", empty_context)
    missing = await csv_rag_tool.query_file_content("q", "missing", ctx)
    assert missing.startswith("Error: File ID")

    monkeypatch.setattr(csv_rag_tool, "get_session_context", context_without_collection)
    no_collection = await csv_rag_tool.query_file_content("q", "f1", ctx)
    assert no_collection.startswith("Error: Collection name not found")

    monkeypatch.setattr(csv_rag_tool, "get_session_context", context_with_collection)
    embed_empty = await csv_rag_tool.query_file_content("q", "f1", ctx)
    assert "embedding failure" in embed_empty

    csv_rag_tool._llm_client_instance = RaisingLLMClient()
    embed_error = await csv_rag_tool.query_file_content("q", "f1", ctx)
    assert embed_error.startswith("Error: Could not process query due to embedding failure")


@pytest.mark.unit
@pytest.mark.asyncio
async def test_query_file_content_success_and_search_failures(monkeypatch: pytest.MonkeyPatch) -> None:
    ctx = FakeCtx()
    llm_client = FakeLLMClient(embeddings=[[0.5]], response_text="final answer")
    csv_rag_tool._llm_client_instance = llm_client
    fake_vsm = FakeVectorStoreManager()

    async def fake_session(ctx: Any, override_session_id: str | None = None) -> str:
        del ctx, override_session_id
        return "session-1"

    async def fake_get_session_context(session_id: str) -> dict[str, Any]:
        return {"uploaded_csv_files": {"f1": {"original_filename": "doc.txt", "collection_name": "col"}}}

    monkeypatch.setattr(csv_rag_tool, "ensure_session_initialized", fake_session)
    monkeypatch.setattr(csv_rag_tool, "get_session_context", fake_get_session_context)
    monkeypatch.setattr(csv_rag_tool, "VectorStoreManager", lambda: fake_vsm)

    answer = await csv_rag_tool.query_file_content("what is this", "f1", ctx)
    assert answer == "final answer"
    assert fake_vsm.query_calls[0]["collection_name"] == "col"
    assert llm_client.response_calls

    fake_vsm.query_result = {"documents": [[]]}
    no_hits = await csv_rag_tool.query_file_content("what is this", "f1", ctx)
    assert no_hits.startswith("Could not find sufficiently relevant information")

    class RaisingVSM(FakeVectorStoreManager):
        def query_collection(self, **kwargs: Any) -> dict[str, Any]:
            raise RuntimeError("vector fail")

    monkeypatch.setattr(csv_rag_tool, "VectorStoreManager", lambda: RaisingVSM())
    search_error = await csv_rag_tool.query_file_content("what is this", "f1", ctx)
    assert search_error.startswith("Error: Could not retrieve relevant data")


@pytest.mark.unit
def test_initialize_nvidia_client_and_register_tools(monkeypatch: pytest.MonkeyPatch) -> None:
    fake_mcp = FakeMCP()
    fake_llm = FakeLLMClient()

    monkeypatch.delenv(csv_rag_tool.NVIDIA_API_KEY_ENV_VAR, raising=False)
    csv_rag_tool._initialize_nvidia_client()
    assert csv_rag_tool._nvidia_async_client is None

    monkeypatch.setenv(csv_rag_tool.NVIDIA_API_KEY_ENV_VAR, "token")
    monkeypatch.setattr(
        csv_rag_tool.openai,
        "AsyncOpenAI",
        lambda base_url, api_key: SimpleNamespace(base_url=base_url, api_key=api_key),
    )
    csv_rag_tool._initialize_nvidia_client()
    assert csv_rag_tool._nvidia_async_client is not None

    csv_rag_tool.register_tools(fake_mcp, fake_llm)
    assert csv_rag_tool._llm_client_instance is fake_llm
    assert fake_mcp.registered == ["process_uploaded_file", "query_file_content"]
